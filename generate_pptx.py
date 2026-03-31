"""포트폴리오 PPT 생성 스크립트"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ── 색상 ──
WHITE = RGBColor(255, 255, 255)
BLACK = RGBColor(0, 0, 0)
DARK = RGBColor(26, 26, 26)
GRAY = RGBColor(100, 100, 100)
LIGHT_GRAY = RGBColor(160, 160, 160)
BG_LIGHT = RGBColor(250, 250, 250)
ACCENT = RGBColor(226, 117, 0)  # #e27500
BLUE = RGBColor(43, 79, 203)
GREEN = RGBColor(22, 163, 74)
SKY = RGBColor(14, 165, 233)
RED_SOFT = RGBColor(220, 80, 60)
AMBER = RGBColor(217, 148, 74)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

def add_bg(slide, color=WHITE):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_shape(slide, left, top, w, h, color, radius=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, left, top, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    if radius:
        shape.adjustments[0] = radius
    return shape

def add_text(slide, left, top, w, h, text, size=14, bold=False, color=DARK, align=PP_ALIGN.LEFT, font_name='맑은 고딕'):
    txBox = slide.shapes.add_textbox(left, top, w, h)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font_name
    p.alignment = align
    return tf

def add_para(tf, text, size=14, bold=False, color=DARK, space_before=0, font_name='맑은 고딕'):
    p = tf.add_paragraph()
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font_name
    p.space_before = Pt(space_before)
    return p


# ════════════════════════════════════════
# SLIDE 1: 표지
# ════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK)

add_text(slide, Inches(1.5), Inches(1.5), Inches(10), Inches(0.6),
         'PORTFOLIO', 16, True, ACCENT, font_name='Consolas')

add_text(slide, Inches(1.5), Inches(2.2), Inches(10), Inches(1.5),
         '윤경은', 52, True, WHITE, font_name='맑은 고딕')

tf = add_text(slide, Inches(1.5), Inches(3.8), Inches(10), Inches(0.8),
              '통계 기반 분석 + 풀스택 개발로 데이터를 서비스로 전환합니다', 20, False, LIGHT_GRAY)

# 구분선
add_shape(slide, Inches(1.5), Inches(5.0), Inches(2), Pt(3), ACCENT)

tf = add_text(slide, Inches(1.5), Inches(5.3), Inches(10), Inches(1.5),
              '', 13, False, LIGHT_GRAY)
add_para(tf, '가천대학교 응용통계학과 4학년', 13, False, LIGHT_GRAY, 0)
add_para(tf, 'yge0307@gmail.com  |  github.com/ykgstar37-lab', 13, False, LIGHT_GRAY, 6)
add_para(tf, 'portfolio-nine-orcin-86.vercel.app', 13, False, ACCENT, 6)


# ════════════════════════════════════════
# SLIDE 2: About Me
# ════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)

add_text(slide, Inches(1), Inches(0.5), Inches(3), Inches(0.4),
         'ABOUT ME', 12, True, ACCENT, font_name='Consolas')
add_text(slide, Inches(1), Inches(0.9), Inches(11), Inches(0.7),
         '어떤 사람인가', 32, True, DARK)

# 왼쪽: 핵심 요약
box = add_shape(slide, Inches(1), Inches(2.0), Inches(5.5), Inches(4.5), BG_LIGHT, 0.05)
tf = add_text(slide, Inches(1.4), Inches(2.3), Inches(4.8), Inches(4.0),
              '', 14, False, DARK)
for line in [
    '• 가천대학교 응용통계학과 4학년',
    '',
    '• 통계 분석(GARCH, Monte Carlo, K-means)을',
    '   풀스택 웹 서비스로 전환하는 데 집중',
    '',
    '• Jupyter/R 분석 → 실시간 대시보드,',
    '   인터랙티브 지도 등 사용자 서비스로 확장',
    '',
    '• LangChain/LangGraph 기반',
    '   AI 에이전트 및 RAG 파이프라인 구축',
    '',
    '• Python(FastAPI, Django) +',
    '   JavaScript/TypeScript(React) 풀스택',
]:
    add_para(tf, line, 13, False, DARK if line.startswith('•') or line.startswith('   ') else DARK, 2)

# 오른쪽: 기술 스택
add_text(slide, Inches(7.2), Inches(2.0), Inches(5), Inches(0.4),
         'TECH STACK', 11, True, ACCENT, font_name='Consolas')

stack_items = [
    ('Language', 'Python, JavaScript, TypeScript'),
    ('Frontend', 'React, Tailwind CSS, Vite, Zustand'),
    ('Backend', 'FastAPI, Django'),
    ('Data / AI', 'pandas, scikit-learn, LangChain, OpenAI'),
    ('Database', 'PostgreSQL, SQLite, Qdrant'),
    ('Infra', 'AWS, Vercel, Render, Docker'),
]

y = 2.5
for cat, items in stack_items:
    add_text(slide, Inches(7.2), Inches(y), Inches(5), Inches(0.3),
             cat, 11, True, GRAY, font_name='Consolas')
    add_text(slide, Inches(7.2), Inches(y + 0.3), Inches(5), Inches(0.3),
             items, 12, False, DARK)
    y += 0.75


# ════════════════════════════════════════
# SLIDE 3: 프로젝트 Overview
# ════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)

add_text(slide, Inches(1), Inches(0.5), Inches(3), Inches(0.4),
         'PROJECTS', 12, True, ACCENT, font_name='Consolas')
add_text(slide, Inches(1), Inches(0.9), Inches(11), Inches(0.7),
         '프로젝트 요약', 32, True, DARK)

projects = [
    ('01', 'CryptoVol Dashboard', 'Personal', '실시간 멀티코인 변동성 예측 대시보드', 'FastAPI · React · WebSocket · GARCH', BLUE),
    ('02', 'Seoul Culture Map', 'Personal', '서울 2,500+ 문화시설 인터랙티브 탐색 맵', 'FastAPI · React · Leaflet · scikit-learn', SKY),
    ('03', 'WorkFlow Agent (DUDE)', 'Team', 'Multi-Agent 사내 업무 자동화 시스템', 'LangGraph · vLLM · LoRA · Qdrant', ACCENT),
    ('04', 'PyMate', 'Team', 'AI 기반 부트캠프 학습 플랫폼', 'Django · LangChain · Qdrant · AWS', GREEN),
    ('05', 'SubFlow', 'Team', '구독 관리 플랫폼 (지출 분석·환율 추적)', 'FastAPI · React · TypeScript · PostgreSQL', RGBColor(139, 92, 246)),
]

y_start = 2.0
for i, (num, name, badge, desc, tech, color) in enumerate(projects):
    y = y_start + i * 1.0
    # 번호
    circle = add_shape(slide, Inches(1), Inches(y), Inches(0.45), Inches(0.45), color, 0.5)
    circle.text_frame.paragraphs[0].text = num
    circle.text_frame.paragraphs[0].font.size = Pt(11)
    circle.text_frame.paragraphs[0].font.bold = True
    circle.text_frame.paragraphs[0].font.color.rgb = WHITE
    circle.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    circle.text_frame.paragraphs[0].font.name = 'Consolas'

    # 프로젝트명
    add_text(slide, Inches(1.7), Inches(y - 0.05), Inches(3.5), Inches(0.35),
             name, 15, True, DARK)
    # 배지
    badge_color = BLUE if badge == 'Personal' else AMBER
    b = add_shape(slide, Inches(5.2), Inches(y + 0.02), Inches(0.8), Inches(0.3), badge_color, 0.5)
    b.text_frame.paragraphs[0].text = badge
    b.text_frame.paragraphs[0].font.size = Pt(8)
    b.text_frame.paragraphs[0].font.bold = True
    b.text_frame.paragraphs[0].font.color.rgb = WHITE
    b.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    b.text_frame.paragraphs[0].font.name = 'Consolas'

    # 설명
    add_text(slide, Inches(6.3), Inches(y - 0.05), Inches(4), Inches(0.3),
             desc, 12, False, GRAY)
    # 기술
    add_text(slide, Inches(6.3), Inches(y + 0.25), Inches(4.5), Inches(0.3),
             tech, 10, False, LIGHT_GRAY, font_name='Consolas')


# ════════════════════════════════════════
# SLIDE 4-6: Problem Solving (3 사례)
# ════════════════════════════════════════

cases = [
    {
        'num': '01',
        'project': 'CryptoVol Dashboard',
        'color': BLUE,
        'title': '분석 코드를 실시간 API로 전환하며\n서비스 관점의 설계를 체득',
        'problem': 'Jupyter에서 수동 실행하던 GARCH 모형을 API로 서빙하려니,\n적합(fit) 연산이 수백ms — 동시 요청 시 응답 지연 발생.\n프론트에서 Binance WebSocket 직접 연결 시 CORS + API 키 노출.',
        'solutions': [
            '5분 TTL 인메모리 캐싱 + 120일 윈도우 제한으로 GARCH 재계산 방지\n개별 모형 실패 시 0.0 반환으로 에러 격리',
            'FastAPI WebSocket 릴레이 서버 구현\nSet 기반 클라이언트 추적으로 Binance → 브라우저 브로드캐스트',
            '코인 전환 시 Promise.all로 API 호출 병렬화\n체감 전환 속도 확보',
        ],
        'results': [('Endpoints', '14 REST + 1 WS'), ('Coins', 'BTC / ETH / SOL'), ('Deploy', 'Render + Vercel')],
        'insight': '"분석할 수 있다"와 "API로 서빙할 수 있다"는\n완전히 다른 설계 관점이 필요하다는 것을 체감',
    },
    {
        'num': '02',
        'project': 'PyMate',
        'color': GREEN,
        'title': 'Flask MVP → Django 프로덕션 전환과\n인프라 구조 이해',
        'problem': 'Flask MVP는 기능적으로 동작했지만, ORM 마이그레이션·정적 파일\n서빙·관리자 페이지 등 프로덕션 기능을 수동 구성해야 했음.\nAWS 배포 시 Nginx → Gunicorn → Flask에서 반복적인 502 에러.',
        'solutions': [
            'Django 내장 기능(ORM migration, admin, collectstatic)으로\n프로덕션 인프라 표준화',
            'RAG 임베딩 768D → 3072D 교체 + Qdrant 벡터 DB 재설계',
            'Nginx → Gunicorn → Django 서버 흐름 직접 구성\n소켓 바인딩 설정 문제 해결로 502 에러 근본 해결',
        ],
        'results': [('Context Precision', '0.83 → 0.97'), ('Context Recall', '0.70 → 0.79'), ('Infra', 'AWS 안정화')],
        'insight': '"기능을 더 만드는 것"보다 "구조를 바꾸는 것"이\n서비스 품질에 더 큰 영향을 준다는 판단 기준을 획득',
    },
    {
        'num': '03',
        'project': 'WorkFlow Agent',
        'color': ACCENT,
        'title': 'LLM 출력을 맹신하지 않는\n서버 검증 시스템 설계',
        'problem': 'sLLM이 confidence 0.95를 출력해도 존재하지 않는 조항을\n인용(환각)하거나 오답 반환. 기본 판단 정확도 37.2%.\nLoRA v2에서 98건 추가 시 오히려 -3.2%p 하락 (라벨 오염).',
        'solutions': [
            '4중 Guardrail: 키워드 매칭(0~1.0), 조항 존재 검증,\n카테고리 제한, 일관성 모니터링(500건 FIFO)',
            '5-factor Confidence 보정 공식 설계\nHard Cap으로 LLM 과신 방지 (RAG < 0.2 → max 0.4)',
            'LoRA v3 — 19건만 약점(재량 표현 14 + 경계 5)\n정밀 타겟팅하여 회복. 데이터 양 < 질 증명',
        ],
        'results': [('판단 정확도', '37% → 85%'), ('JSON 유효율', '70% → 97%'), ('RAG MRR', '0.63 → 0.95')],
        'insight': '"데이터 양 < 데이터 질" — 98건 무작위 추가보다\n19건 정밀 타겟팅이 효과적. LLM 서비스에서 서버 단 검증이 핵심',
    },
]

for case in cases:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(slide, WHITE)

    # 헤더
    add_text(slide, Inches(1), Inches(0.4), Inches(5), Inches(0.3),
             f'PROBLEM SOLVING  #{case["num"]}', 11, True, case['color'], font_name='Consolas')
    add_text(slide, Inches(1), Inches(0.7), Inches(6), Inches(0.3),
             case['project'], 11, True, GRAY)
    add_text(slide, Inches(1), Inches(1.0), Inches(10), Inches(1.0),
             case['title'], 26, True, DARK)

    # Problem
    prob_box = add_shape(slide, Inches(1), Inches(2.3), Inches(5.5), Inches(2.0), RGBColor(254, 242, 242), 0.03)
    add_text(slide, Inches(1.3), Inches(2.4), Inches(1.5), Inches(0.3),
             'PROBLEM', 10, True, RED_SOFT, font_name='Consolas')
    add_text(slide, Inches(1.3), Inches(2.8), Inches(4.9), Inches(1.4),
             case['problem'], 11, False, RGBColor(120, 50, 40))

    # Solutions
    add_text(slide, Inches(7), Inches(2.3), Inches(2), Inches(0.3),
             'SOLUTION', 10, True, case['color'], font_name='Consolas')

    for j, sol in enumerate(case['solutions']):
        y = 2.7 + j * 1.05
        circle = add_shape(slide, Inches(7), Inches(y), Inches(0.3), Inches(0.3), case['color'], 0.5)
        circle.text_frame.paragraphs[0].text = str(j + 1)
        circle.text_frame.paragraphs[0].font.size = Pt(9)
        circle.text_frame.paragraphs[0].font.bold = True
        circle.text_frame.paragraphs[0].font.color.rgb = WHITE
        circle.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        circle.text_frame.paragraphs[0].font.name = 'Consolas'
        add_text(slide, Inches(7.5), Inches(y - 0.05), Inches(4.8), Inches(0.9),
                 sol, 10.5, False, DARK)

    # Results
    add_text(slide, Inches(1), Inches(4.6), Inches(2), Inches(0.3),
             'RESULT', 10, True, GREEN, font_name='Consolas')

    for k, (label, value) in enumerate(case['results']):
        x = 1 + k * 1.9
        r_box = add_shape(slide, Inches(x), Inches(4.95), Inches(1.7), Inches(0.9), BG_LIGHT, 0.05)
        add_text(slide, Inches(x + 0.1), Inches(5.0), Inches(1.5), Inches(0.25),
                 label, 9, True, GRAY, PP_ALIGN.CENTER, 'Consolas')
        add_text(slide, Inches(x + 0.1), Inches(5.3), Inches(1.5), Inches(0.35),
                 value, 13, True, DARK, PP_ALIGN.CENTER)

    # Insight
    insight_box = add_shape(slide, Inches(7), Inches(5.8), Inches(5.3), Inches(1.2), RGBColor(245, 245, 245), 0.03)
    add_text(slide, Inches(7.3), Inches(5.85), Inches(2), Inches(0.3),
             'KEY INSIGHT', 9, True, case['color'], font_name='Consolas')
    add_text(slide, Inches(7.3), Inches(6.2), Inches(4.7), Inches(0.7),
             case['insight'], 11, True, RGBColor(60, 60, 60))


# ════════════════════════════════════════
# SLIDE 7: 마무리
# ════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, DARK)

add_text(slide, Inches(1.5), Inches(2.0), Inches(10), Inches(1.0),
         'Thank You', 48, True, WHITE, PP_ALIGN.CENTER, 'Consolas')

add_shape(slide, Inches(5.5), Inches(3.3), Inches(2.3), Pt(3), ACCENT)

tf = add_text(slide, Inches(1.5), Inches(3.8), Inches(10), Inches(2.5),
              '', 14, False, LIGHT_GRAY, PP_ALIGN.CENTER)
add_para(tf, '윤경은  |  가천대학교 응용통계학과', 15, False, WHITE, 0)
add_para(tf, '', 10, False, WHITE, 8)
add_para(tf, 'yge0307@gmail.com', 13, False, LIGHT_GRAY, 4)
add_para(tf, 'github.com/ykgstar37-lab', 13, False, LIGHT_GRAY, 4)
add_para(tf, 'linkedin.com/in/yoon0307eun', 13, False, LIGHT_GRAY, 4)
add_para(tf, '', 10, False, WHITE, 12)
add_para(tf, 'portfolio-nine-orcin-86.vercel.app', 14, True, ACCENT, 4)


# ── 저장 ──
output = r'c:\Users\ykgst\OneDrive\바탕 화면\portfolio\Portfolio_윤경은.pptx'
prs.save(output)
print(f'Done: {output}')
